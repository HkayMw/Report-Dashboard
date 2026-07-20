"""
Generates a formatted .xlsx export: cleaned data, summary, per-zone,
per-center sheets, with native Excel charts.

Model note: Birth Registration and NID (National ID) Registration are two
SEPARATE activities captured on the same daily form - not a before/after
pipeline. They are reported side-by-side throughout, never as a "rate".
"""
import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, Reference

from data_processing import NUMERIC_COLS

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

EXPORT_COLS = [
    "Timestamp", "ZONE NAME", "CENTER NAME", "Date",
    "Total Male Births Registered", "Total Female Births Registered",
    "Total Males Processed", "Total Females Processed",
    "Total Birth Registrations", "Total NID Registrations",
]


def _write_df(ws, df, start_row=1, start_col=1, index=False):
    """Write a dataframe to a worksheet starting at (start_row, start_col),
    styled header. start_col > 1 lets several tables sit side by side."""
    cols = ([df.index.name or ""] if index else []) + list(df.columns)
    for j, col_name in enumerate(cols, start=start_col):
        cell = ws.cell(row=start_row, column=j, value=str(col_name))
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for i, (idx, row) in enumerate(df.iterrows(), start=start_row + 1):
        offset = start_col
        if index:
            ws.cell(row=i, column=start_col, value=idx)
            offset = start_col + 1
        for j, val in enumerate(row, start=offset):
            if pd.isna(val):
                val = None
            elif hasattr(val, "isoformat"):
                val = val.isoformat()
            ws.cell(row=i, column=j, value=val)

    for j, col_name in enumerate(cols, start=start_col):
        rel = j - start_col
        max_len = max(
            [len(str(col_name))] + [len(str(v)) for v in (df.index if index and rel == 0 else df.iloc[:, rel - (1 if index else 0)])]
        ) if len(df) else len(str(col_name))
        ws.column_dimensions[get_column_letter(j)].width = min(max(max_len + 2, 10), 40)

    return start_row + len(df) + 1  # next free row


def _districts_in(df: pd.DataFrame) -> list:
    """Distinct non-empty district labels in a summary frame ([] when the
    frame has no District column)."""
    if "District" not in df.columns:
        return []
    return sorted(d for d in df["District"].astype(str).unique() if d.strip())


def _add_bar_chart(ws, title, y_title, x_title, value_col, label_col,
                   header_row, n_rows, anchor, size):
    chart = BarChart()
    chart.title = title
    chart.y_axis.title = y_title
    chart.x_axis.title = x_title
    data = Reference(ws, min_col=value_col, min_row=header_row,
                     max_row=header_row + n_rows, max_col=value_col)
    cats = Reference(ws, min_col=label_col, min_row=header_row + 1,
                     max_row=header_row + n_rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width, chart.height = size
    ws.add_chart(chart, anchor)


# One anchored chart occupies roughly this many worksheet rows at the chart
# heights used here — spacing constant for stacking charts above tables.
_CHART_ROWS = 21


def _write_summary_sheet(ws, df, label_col, entity, chart_size, chart_top_n=None):
    """Write a per-zone/per-center summary sheet with GRAPHS ON TOP and the
    data below them. With multiple districts, each district gets its own
    vertical band — its charts stacked at the top, its table underneath —
    and the bands sit side by side for easy district comparison."""
    districts = _districts_in(df)
    if len(districts) <= 1:
        n = len(df)
        table_start = (_CHART_ROWS + 2) if n else 1
        _write_df(ws, df, start_row=table_start)
        if n:
            n_chart = min(n, chart_top_n or n)
            suffix = f" (top {n_chart})" if n_chart < n else ""
            col_birth = list(df.columns).index("Total Birth Registrations") + 1
            col_nid = list(df.columns).index("Total NID Registrations") + 1
            _add_bar_chart(ws, f"Birth Registrations by {entity}{suffix}", "Births Registered",
                           entity, col_birth, 1, table_start, n_chart, "A1", chart_size)
            _add_bar_chart(ws, f"NID Registrations by {entity}{suffix}", "NID Registrations",
                           entity, col_nid, 1, table_start, n_chart, "K1", chart_size)
        return

    n_cols = len(df.columns)
    band_width = n_cols + 1  # one spacer column between district bands
    # Charts must fit inside their band so neighbouring districts' charts
    # don't overlap: ~1.8cm per default-width column.
    band_chart_size = (max(10, int(n_cols * 1.8)), chart_size[1])
    table_start = 2 + 2 * _CHART_ROWS  # title row + two stacked charts

    for i, d in enumerate(districts):
        band_col = 1 + i * band_width
        col_letter = get_column_letter(band_col)
        sub = df[df["District"].astype(str) == d]
        n = len(sub)
        ws.cell(row=1, column=band_col, value=d).font = Font(bold=True, size=14)
        _write_df(ws, sub, start_row=table_start, start_col=band_col)
        if not n:
            continue
        n_chart = min(n, chart_top_n or n)
        cols = list(sub.columns)
        col_birth = band_col + cols.index("Total Birth Registrations")
        col_nid = band_col + cols.index("Total NID Registrations")
        label_idx = band_col + cols.index(label_col)
        suffix = f" (top {n_chart})" if n_chart < n else ""
        _add_bar_chart(ws, f"Birth Registrations by {entity} — {d}{suffix}", "Births Registered",
                       entity, col_birth, label_idx, table_start, n_chart,
                       f"{col_letter}2", band_chart_size)
        _add_bar_chart(ws, f"NID Registrations by {entity} — {d}{suffix}", "NID Registrations",
                       entity, col_nid, label_idx, table_start, n_chart,
                       f"{col_letter}{2 + _CHART_ROWS}", band_chart_size)


def build_report(cleaned_df: pd.DataFrame, report: dict, zone_summary: pd.DataFrame,
                  center_summary: pd.DataFrame, kpis: dict) -> bytes:
    wb = Workbook()

    # ---------- Summary sheet ----------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "NID & NRBC Daily Reporting — Summary"
    ws["A1"].font = Font(bold=True, size=14)

    row = 3
    for label, value in kpis.items():
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    # Per-district totals — always shown when more than one district is in
    # view (i.e. the ALL scope), so districts are never lumped together.
    zs_all = zone_summary.reset_index()
    _summary_districts = _districts_in(zs_all)
    if len(_summary_districts) > 1:
        row += 1
        ws.cell(row=row, column=1, value="Totals by District").font = Font(bold=True, size=12)
        row += 1
        num_cols = [c for c in zs_all.columns if c not in ("ZONE NAME", "District")]
        per_district = zs_all.groupby("District")[num_cols].sum().reset_index()
        row = _write_df(ws, per_district, start_row=row) + 1

    row += 1
    ws.cell(row=row, column=1, value="Data Quality Notes").font = Font(bold=True, size=12)
    row += 1
    notes = [
        f"Duplicate/resubmitted rows removed (kept latest by timestamp): {report.get('duplicates_removed_count', 0)}",
        f"Rows with missing numeric values (filled with 0): {len(report.get('missing_numeric_rows', []))}",
        f"Rows with invalid/negative numbers: {len(report.get('negative_rows', []))}",
        f"Rows with non-whole-number values (likely data entry errors, e.g. 0.01): {len(report.get('non_integer_rows', []))}",
    ]
    for note in notes:
        ws.cell(row=row, column=1, value=note)
        row += 1

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 20

    # ---------- Cleaned Data sheet ----------
    ws2 = wb.create_sheet("Cleaned Data")
    data_cols = EXPORT_COLS
    data_sheet_df = cleaned_df
    if "District" in cleaned_df.columns:
        data_cols = ["District"] + EXPORT_COLS
        data_sheet_df = cleaned_df.sort_values(["District", "Date"])
    _write_df(ws2, data_sheet_df[data_cols])
    ws2.freeze_panes = "A2"

    # ---------- Per-Zone sheet ----------
    ws3 = wb.create_sheet("Per-Zone")
    _write_summary_sheet(
        ws3, zone_summary.reset_index(), label_col="ZONE NAME", entity="Zone",
        chart_size=(20, 10),
    )

    # ---------- Per-Center sheet ----------
    ws4 = wb.create_sheet("Per-Center")
    _write_summary_sheet(
        ws4, center_summary.reset_index(), label_col="CENTER NAME", entity="Center",
        chart_size=(24, 12), chart_top_n=15,
    )

    # ---------- Trend sheet (by date) ----------
    # Separate trends per ACTIVITY (Birth vs NID — they're independent
    # activities, never merged) and per DISTRICT (one line per district when
    # several are in view, plus an overall total). Charts on top, data below.
    ws5 = wb.create_sheet("Trend")
    day = cleaned_df["Date"].dt.date
    activities = [
        ("Birth Registrations", "Total Birth Registrations"),
        ("NID Registrations", "Total NID Registrations"),
    ]
    trend_districts = (
        _districts_in(cleaned_df) if "District" in cleaned_df.columns else []
    )
    trend_row = _CHART_ROWS + 2
    for i, (label, value_col) in enumerate(activities):
        if len(trend_districts) > 1:
            pivot = (
                cleaned_df.groupby([day, "District"])[value_col].sum()
                .unstack(fill_value=0)
            )
            pivot["All Districts"] = pivot.sum(axis=1)
        else:
            pivot = cleaned_df.groupby(day)[[value_col]].sum()
            pivot.columns = [label]
        pivot.index.name = "Date"
        tbl = pivot.reset_index()

        ws5.cell(row=trend_row, column=1, value=f"{label} per day").font = Font(bold=True, size=12)
        header_row = trend_row + 1
        next_free = _write_df(ws5, tbl, start_row=header_row)
        n_days = len(tbl)
        if n_days > 1:
            chart = LineChart()
            chart.title = f"{label} Over Time" + (" — by District" if len(trend_districts) > 1 else "")
            chart.y_axis.title = label
            chart.x_axis.title = "Date"
            data = Reference(ws5, min_col=2, min_row=header_row,
                             max_row=header_row + n_days, max_col=len(tbl.columns))
            cats = Reference(ws5, min_col=1, min_row=header_row + 1, max_row=header_row + n_days)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.width, chart.height = 16, 10
            # Charts sit side by side at the top; their tables follow below.
            ws5.add_chart(chart, f"{get_column_letter(1 + i * 10)}1")
        trend_row = next_free + 1

    # ---------- Flagged rows sheet ----------
    ws6 = wb.create_sheet("Flagged Rows")
    r = 1
    ws6.cell(row=r, column=1, value="Duplicate/Resubmitted Rows Found").font = Font(bold=True, size=12)
    r += 1
    dupes = report.get("duplicates_found", pd.DataFrame())
    if len(dupes):
        r = _write_df(ws6, dupes[["Timestamp", "ZONE NAME", "CENTER NAME", "Date"] + NUMERIC_COLS], start_row=r) + 1
    else:
        ws6.cell(row=r, column=1, value="None found")
        r += 2

    ws6.cell(row=r, column=1, value="Rows with Missing Numeric Values").font = Font(bold=True, size=12)
    r += 1
    missing = report.get("missing_numeric_rows", pd.DataFrame())
    if len(missing):
        cols_present = [c for c in EXPORT_COLS if c in missing.columns]
        r = _write_df(ws6, missing[cols_present], start_row=r) + 1
    else:
        ws6.cell(row=r, column=1, value="None found")
        r += 2

    ws6.cell(row=r, column=1, value="Rows with Non-Whole-Number Values (likely data entry errors)").font = Font(bold=True, size=12)
    r += 1
    non_int = report.get("non_integer_rows", pd.DataFrame())
    if len(non_int):
        cols_present = [c for c in EXPORT_COLS if c in non_int.columns]
        r = _write_df(ws6, non_int[cols_present], start_row=r) + 1
    else:
        ws6.cell(row=r, column=1, value="None found")
        r += 2

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_completion_report(district_df: pd.DataFrame, zone_df: pd.DataFrame,
                            center_df: pd.DataFrame, start_date, end_date) -> bytes:
    """Excel export for master_reference.completion_report()'s output: a
    district summary sheet, then per-zone and per-center sheets — figures only
    reflect submissions already mapped to the predefined master list."""
    wb = Workbook()

    ws0 = wb.active
    ws0.title = "District Completion"
    ws0["A1"] = f"District Completion — {start_date} to {end_date}"
    ws0["A1"].font = Font(bold=True, size=14)
    _write_df(ws0, district_df, start_row=3)

    ws = wb.create_sheet("Zone Completion")
    ws["A1"] = f"Zone Completion — {start_date} to {end_date}"
    ws["A1"].font = Font(bold=True, size=14)
    _write_df(ws, zone_df, start_row=3)

    ws2 = wb.create_sheet("Center Completion")
    ws2["A1"] = f"Center Completion — {start_date} to {end_date}"
    ws2["A1"].font = Font(bold=True, size=14)
    _write_df(ws2, center_df, start_row=3)
    ws2.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
